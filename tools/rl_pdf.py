import openpyxl, sys
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle
from reportlab.lib import colors
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
PS=ParagraphStyle('p',fontName='Helvetica',fontSize=6.6,leading=7.6)
src, out = sys.argv[1], sys.argv[2]
ws = openpyxl.load_workbook(src).active
rows = [list(r) for r in ws.iter_rows(min_row=16, values_only=True) if r[0] is not None]
spans = [m for m in ws.merged_cells.ranges if m.min_row >= 16]
s = lambda v: '' if v is None else str(v).strip()
hdr = ['#','Full Name','Birthdate','Gender','Passport No.','Passport Exp.','Room','Group','Type','Special Request','Free Gift']
data = [hdr] + [[Paragraph(s(v),PS) if j in (9,10) and s(v) else s(v) for j,v in enumerate(r[:11])] for r in rows]
W, H = landscape(A4)
c = canvas.Canvas(out, pagesize=(W, H), pageCompression=1)
RED = colors.HexColor('#BE1E2D')
c.setFillColor(RED); c.rect(0, H-46, W, 46, fill=1, stroke=0)
c.setFillColor(colors.white); c.setFont('Helvetica-Bold', 14)
c.drawString(22, H-28, 'ARBA Travel  -  Room List  SEL/JT/STD/031126/CX')
c.setFont('Helvetica', 8.5)
c.drawRightString(W-22, H-20, 'Seoul  3 - 8 Nov 2026  |  Final (Airtable, 28 Oct 2026)')
c.drawRightString(W-22, H-33, 'TL 1 / Adult 25 / Child 2 / Child NB 2 / Infant 1 = 31   |   Rooms: SINGLE 2 / TWIN 4 / DOUBLE 3 / TRIPLE 4 (13)')
t = Table(data, colWidths=[16,160,50,36,52,52,40,156,28,104,104], rowHeights=[14]+[14.6]*len(rows))
st = [('FONT',(0,0),(-1,-1),'Helvetica',6.6), ('FONT',(0,0),(-1,0),'Helvetica-Bold',6.8),
      ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#3A3130')), ('TEXTCOLOR',(0,0),(-1,0),colors.white),
      ('GRID',(0,0),(-1,-1),0.25,colors.HexColor('#BBBBBB')), ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
      ('LEFTPADDING',(0,0),(-1,-1),2.5), ('RIGHTPADDING',(0,0),(-1,-1),2), ('TOPPADDING',(0,0),(-1,-1),0), ('BOTTOMPADDING',(0,0),(-1,-1),1),
      ('ALIGN',(0,0),(0,-1),'CENTER'), ('ALIGN',(6,0),(6,-1),'CENTER'), ('ALIGN',(8,0),(8,-1),'CENTER'), ('FONT',(6,1),(6,-1),'Helvetica-Bold',6.6)]
for m in spans:
    st.append(('SPAN',(m.min_col-1, m.min_row-15),(m.max_col-1, m.max_row-15)))
for i, r in enumerate(data[1:], 1):
    if r[6] and r[6] != 'NO BED': st.append(('LINEABOVE',(0,i),(-1,i),1.1,colors.black))
    if r[8] == 'TL': st.append(('BACKGROUND',(0,i),(-1,i),colors.HexColor('#DCD3CF')))
    if r[8] in ('CHD','INF'): st.append(('TEXTCOLOR',(8,i),(8,i),RED))
t.setStyle(TableStyle(st))
w, h = t.wrapOn(c, W, H)
t.drawOn(c, (W-w)/2, H-56-h)
c.setFillColor(colors.HexColor('#6E6461')); c.setFont('Helvetica', 7)
c.drawString(22, 14, 'ARBA Travel Korea Co., Ltd. - Confidential (passport data)')
c.save()
print(w, h)
